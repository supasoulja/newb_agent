"""
Document RAG — ingest, chunk, embed, and search uploaded documents.

Supported formats: .txt .md .pdf .docx .csv .py .json
Chunks are 800 chars with 100-char overlap. Each chunk is embedded with
the fast CPU ONNX model (384-dim) and stored in a sqlite-vec vec0 table.
At shutdown, chunks are re-embedded with qwen3-embedding:4b (2560-dim)
into a shadow HQ table.

At query time, context.py auto-injects the top-K most relevant chunks.
The docs.search tool lets the model do targeted retrieval.
"""
from __future__ import annotations
import uuid
from datetime import datetime
from pathlib import Path
from typing import Callable

from kai.config import DEBUG
from kai.db import get_conn, sqlite_vec_available

EmbedFn = Callable[[str], list[float]]

CHUNK_SIZE    = 800   # characters per chunk
CHUNK_OVERLAP = 100   # overlap between consecutive chunks
ALLOWED_TYPES = {".txt", ".md", ".pdf", ".docx", ".csv", ".py", ".json"}


# ── Text extraction ────────────────────────────────────────────────────────────

def _extract_text(filepath: Path, original_name: str) -> str:
    """Extract plain text from a file. Uses original_name for type detection."""
    suffix = Path(original_name).suffix.lower()

    if suffix in {".txt", ".md", ".py", ".json", ".csv"}:
        try:
            return filepath.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            raise ValueError(f"Could not read file: {e}")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(filepath))
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                pages.append(text)
            return "\n\n".join(pages)
        except Exception as e:
            raise ValueError(f"Could not read PDF: {e}")

    if suffix in {".docx", ".doc"}:
        try:
            import docx
            doc = docx.Document(str(filepath))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            raise ValueError(f"Could not read Word document: {e}")

    raise ValueError(f"Unsupported file type: {suffix!r}")


# ── Chunking ───────────────────────────────────────────────────────────────────

def _chunk(text: str) -> list[str]:
    """Split text into overlapping chunks."""
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - CHUNK_OVERLAP
    return chunks


# ── Public API ─────────────────────────────────────────────────────────────────

def ingest(
    filepath: Path,
    embed_fn: EmbedFn,
    original_name: str | None = None,
    user_id: int = 0,
) -> dict:
    """
    Extract text from filepath, chunk it, embed each chunk, store everything.
    Returns metadata dict: {doc_id, filename, file_type, char_count, chunk_count}.

    original_name: the browser filename (used for type detection + display).
    If omitted, filepath.name is used.
    """
    filename = original_name or filepath.name
    suffix   = Path(filename).suffix.lower()

    if suffix not in ALLOWED_TYPES:
        raise ValueError(f"Unsupported type {suffix!r}. Allowed: {', '.join(sorted(ALLOWED_TYPES))}")

    text = _extract_text(filepath, filename)
    if not text.strip():
        raise ValueError("Document appears to be empty or unreadable.")

    chunks = _chunk(text)
    if not chunks:
        raise ValueError("No text could be extracted from the document.")

    doc_id     = str(uuid.uuid4())
    now        = datetime.now().isoformat()
    file_type  = suffix.lstrip(".")
    char_count = len(text)

    # Persist document record and all chunks in one transaction
    conn = get_conn()
    conn.execute(
        "INSERT INTO rag_documents (doc_id, user_id, filename, file_type, char_count, chunk_count, uploaded_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (doc_id, user_id, filename, file_type, char_count, len(chunks), now),
    )
    chunk_rows = []
    for i, chunk_text in enumerate(chunks):
        chunk_id = str(uuid.uuid4())
        chunk_rows.append((chunk_id, doc_id, i, chunk_text))
    conn.executemany(
        "INSERT INTO rag_chunks (chunk_id, doc_id, chunk_index, content) VALUES (?, ?, ?, ?)",
        chunk_rows,
    )
    conn.commit()

    # Embed chunks — best-effort, text is already committed above
    if sqlite_vec_available():
        try:
            import sqlite_vec
            # Fetch rowids for the chunks we just inserted
            rows = conn.execute(
                "SELECT rowid, chunk_id FROM rag_chunks WHERE doc_id = ? ORDER BY chunk_index",
                (doc_id,),
            ).fetchall()

            # Embed in one batch call for speed (CPU, no VRAM)
            texts = list(chunks)
            from kai.embed import embed_batch as _fast_embed_batch
            embeddings = _fast_embed_batch(texts)

            for (rowid, _chunk_id), emb in zip(rows, embeddings):
                conn.execute(
                    "INSERT INTO rag_chunks_vec (rowid, embedding) VALUES (?, ?)",
                    (rowid, sqlite_vec.serialize_float32(emb)),
                )
            conn.commit()
        except Exception:
            if DEBUG:
                import traceback; traceback.print_exc()

    return {
        "doc_id":      doc_id,
        "filename":    filename,
        "file_type":   file_type,
        "char_count":  char_count,
        "chunk_count": len(chunks),
        "uploaded_at": now,
    }


def search(
    query: str,
    embed_fn: EmbedFn | None = None,
    top_k: int = 5,
    query_embedding: list[float] | None = None,
    user_id: int = 0,
) -> list[dict]:
    """
    Find the most relevant chunks for query.
    Returns list of dicts: {doc_id, doc_name, chunk_index, content, distance}.
    Uses vector search when available, falls back to LIKE substring search.
    Scoped to user's own docs + shared docs.
    """
    if (query_embedding or embed_fn) and sqlite_vec_available():
        return _vector_search(query, embed_fn, top_k, query_embedding, user_id)
    return _text_search(query, top_k, user_id)


def _vector_search(
    query: str, embed_fn: EmbedFn, top_k: int,
    query_embedding: list[float] | None = None,
    user_id: int = 0,
) -> list[dict]:
    import sqlite_vec
    embedding = query_embedding or embed_fn(query)
    conn = get_conn()

    # Step 1: pure KNN — no JOINs inside the MATCH query
    knn_rows = conn.execute(
        "SELECT rowid, distance FROM rag_chunks_vec "
        "WHERE embedding MATCH ? ORDER BY distance LIMIT ?",
        (sqlite_vec.serialize_float32(embedding), int(top_k * 3)),
    ).fetchall()
    if not knn_rows:
        return []
    rowid_to_dist = {r[0]: r[1] for r in knn_rows}
    placeholders  = ",".join("?" * len(rowid_to_dist))
    # Step 2: fetch chunk content + doc name, filtered by user ownership or shared
    rows = conn.execute(
        f"SELECT c.rowid, c.doc_id, c.chunk_index, c.content, d.filename "
        f"FROM rag_chunks c "
        f"JOIN rag_documents d ON d.doc_id = c.doc_id "
        f"WHERE c.rowid IN ({placeholders}) AND (d.user_id = ? OR d.shared = 1) "
        f"LIMIT ?",
        (*rowid_to_dist.keys(), user_id, top_k),
    ).fetchall()

    return [
        {
            "doc_id":      row[1],
            "doc_name":    row[4],
            "chunk_index": row[2],
            "content":     row[3],
            "distance":    rowid_to_dist.get(row[0], 9.0),
        }
        for row in rows
    ]


def _text_search(query: str, top_k: int, user_id: int = 0) -> list[dict]:
    conn = get_conn()
    escaped = query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    rows = conn.execute(
        "SELECT c.doc_id, c.chunk_index, c.content, d.filename "
        "FROM rag_chunks c "
        "JOIN rag_documents d ON d.doc_id = c.doc_id "
        "WHERE (d.user_id = ? OR d.shared = 1) AND c.content LIKE ? ESCAPE '\\' "
        "LIMIT ?",
        (user_id, f"%{escaped}%", top_k),
    ).fetchall()
    return [
        {
            "doc_id":      row[0],
            "doc_name":    row[3],
            "chunk_index": row[1],
            "content":     row[2],
            "distance":    0.0,
        }
        for row in rows
    ]


def list_documents(user_id: int = 0) -> list[dict]:
    """Return documents visible to this user (own + shared), newest first."""
    conn = get_conn()
    rows = conn.execute(
        "SELECT doc_id, filename, file_type, char_count, chunk_count, uploaded_at "
        "FROM rag_documents WHERE user_id = ? OR shared = 1 "
        "ORDER BY uploaded_at DESC",
        (user_id,)
    ).fetchall()
    return [
        {
            "doc_id":      r[0],
            "filename":    r[1],
            "file_type":   r[2],
            "char_count":  r[3],
            "chunk_count": r[4],
            "uploaded_at": r[5],
        }
        for r in rows
    ]


def has_documents(user_id: int = 0) -> bool:
    """Fast check — returns True if any documents are visible to this user."""
    conn = get_conn()
    row = conn.execute(
        "SELECT 1 FROM rag_documents WHERE user_id = ? OR shared = 1 LIMIT 1",
        (user_id,)
    ).fetchone()
    return row is not None


def delete_document(doc_id: str, user_id: int = 0) -> bool:
    """Delete a document and all its chunks + vectors. Owner-only enforcement."""
    conn = get_conn()
    row = conn.execute(
        "SELECT 1 FROM rag_documents WHERE doc_id = ? AND user_id = ?",
        (doc_id, user_id)
    ).fetchone()
    if not row:
        return False
    # Get rowids of chunks before deleting (for vec table cleanup)
    chunk_rowids = [
        r[0] for r in conn.execute(
            "SELECT rowid FROM rag_chunks WHERE doc_id = ?", (doc_id,)
        ).fetchall()
    ]
    conn.execute("DELETE FROM rag_chunks WHERE doc_id = ?",   (doc_id,))
    conn.execute("DELETE FROM rag_documents WHERE doc_id = ?", (doc_id,))

    # Clean up vector table
    if chunk_rowids and sqlite_vec_available():
        try:
            placeholders = ",".join("?" * len(chunk_rowids))
            conn.execute(
                f"DELETE FROM rag_chunks_vec WHERE rowid IN ({placeholders})",
                chunk_rowids,
            )
        except Exception:
            if DEBUG:
                import traceback; traceback.print_exc()

    conn.commit()
    return True
